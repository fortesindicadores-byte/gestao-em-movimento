# Planilha das placas da Rota com manutenção R$/km zero no FT (para validar).
# Uso: na pasta com pq/ (do load.py): python3 placas_xlsx.py  →  out/Placas_manutencao_zerada_FT_2026.xlsx
# O .xlsx NÃO vai para o repositório (placas e remuneração).
from comum import *
import datetime as dt
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment
c=ft('ROTA','Caminhao'); c['vig']=c._ano.astype(str)+'-'+c._mes.astype(str).str.zfill(2)
c['mod']=c['Modelo do Cavalo'].astype(str).str.replace(' ',' ').str.replace(r'\.0$','',regex=True).str.strip()
a=c[c.Ativo==True].copy()
z=a[(a['R$/KM'].fillna(0)==0)&(a['Consumo KM/L']>0)&(a.manutencaoVidaMeses>=a.freeMaintenance)&(a.vig>='2026-01')&(a.vig<='2026-10')].copy()
z['pot']=z.manutencaoBid*z['KM Mês Roadshow']/2
z['q']=z._qz.astype(int).astype(str)+'ª '+z.vig.map(lambda v:['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez'][int(v[5:])-1]+'/'+v[2:4])
k=['Unidade - Nome','_vig','mod']
pay=a[a['R$/KM'].fillna(0)>0].groupby(k).agg(irm=('Placa','nunique'),rk=('R$/KM','median')).reset_index()
z=z.merge(pay,on=k,how='left').sort_values(['vig','_qz'])
def dserial(x):
    try: return (dt.datetime(1899,12,30)+dt.timedelta(days=float(x))).date()
    except: return None
g=z.groupby('Placa')
P=g.agg(uni=('Unidade - Nome','last'),chassi=('chassi','last'),mod=('mod','last'),ent=('Data de Entrada','last'),
        vida=('manutencaoVidaMeses','last'),free=('freeMaintenance','last'),nq=('q','count'),q1=('q','first'),q2=('q','last'),
        bid=('manutencaoBid','median'),km=('KM Mês Roadshow','median'),pot=('pot','sum'),irm=('irm','max'),rk=('rk','median'),
        fora=('manutencaoCompraForaDoBidAutorizada','last')).reset_index()
P['cod']=P.uni.map(FTU)
P['faixa']=pd.cut(P.vida,[-1,12,24,1e9],labels=['Até 12 meses','12 a 24 meses','Mais de 24 meses']).astype(str)
P['ord']=P.faixa.map({'Mais de 24 meses':0,'12 a 24 meses':1,'Até 12 meses':2})
P=P.sort_values(['ord','pot'],ascending=[True,False])
wb=Workbook(); ws=wb.active; ws.title='Placas zeradas'
H=['Faixa de vida','Unidade','Placa','Chassi','Modelo','Data de entrada','Vida (meses)','Free maintenance FT (meses)','Quinzenas zeradas em 2026','Primeira','Última','BID R$/km','Km mês (mediana)','Ao valor do BID, jan–out (R$)','Placas do mesmo modelo e unidade COM contrato','R$/km do contrato delas','Compra fora do BID autorizada','Validação (preencher)','Observação']
ws.append(H)
for r in P.itertuples():
    ws.append([r.faixa,r.cod or r.uni,r.Placa,r.chassi,r.mod,dserial(r.ent),round(float(r.vida),1),float(r.free),int(r.nq),r.q1,r.q2,
               None if pd.isna(r.bid) else float(r.bid),None if pd.isna(r.km) else round(float(r.km)),round(float(r.pot),2),
               None if pd.isna(r.irm) else int(r.irm),None if pd.isna(r.rk) else float(r.rk),
               None if pd.isna(r.fora) else ('Sim' if bool(r.fora) else 'Não'),None,None])
n=len(P)+1
fmt={'F':'dd/mm/yyyy','G':'0.0','H':'0','L':'#,##0.00','M':'#,##0','N':'#,##0.00','P':'#,##0.00'}
for col,f in fmt.items():
    for i in range(2,n+1): ws[f'{col}{i}'].number_format=f
t=Table(displayName='PlacasZeradas',ref=f'A1:S{n}'); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
ws.freeze_panes='D2'
for col,w in zip('ABCDEFGHIJKLMNOPQRS',[17,9,10,20,22,12,9,11,11,11,11,9,11,15,14,11,12,22,40]): ws.column_dimensions[col].width=w
for cell in ws[1]: cell.alignment=Alignment(wrap_text=True,vertical='top')
ws.row_dimensions[1].height=48
# detalhe
ws2=wb.create_sheet('Por quinzena')
ws2.append(['Unidade','Placa','Modelo','Quinzena','Vigência FT','Vida (meses)','Contrato R$/km','BID R$/km','Km mês','Ao valor do BID (R$)','Mesmo modelo com contrato','R$/km delas'])
for d in z.to_dict('records'):
    nn=lambda v: None if pd.isna(v) else float(v)
    ws2.append([FTU.get(d['Unidade - Nome']),d['Placa'],d['mod'],d['q'],d['_vig'],round(float(d['manutencaoVidaMeses']),1),
                nn(d['manutencaoContrato']),nn(d['manutencaoBid']),nn(d['KM Mês Roadshow']),round(float(d['pot']),2),
                None if pd.isna(d['irm']) else int(d['irm']),nn(d['rk'])])
m=len(z)+1
t2=Table(displayName='PorQuinzena',ref=f'A1:L{m}'); t2.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws2.add_table(t2); ws2.freeze_panes='C2'
for col,w in zip('ABCDEFGHIJKL',[9,10,22,11,16,9,10,9,10,13,12,10]): ws2.column_dimensions[col].width=w
ws3=wb.create_sheet('Leia-me')
for l in ['Placas da Rota ATIVAS no Freightech, em 2026 (jan–out), com manutenção R$/km = 0 depois do free maintenance do FT (6 meses).',
 'Fonte: robô do Freightech, tabela Caminhão do canal ROTA, uma linha por placa e quinzena.',
 'Vida (meses) = campo manutencaoVidaMeses do FT: sobe 0,5 por quinzena e 1 por mês, e bate com a Data de entrada. É em MESES.',
 'Ao valor do BID = BID R$/km × km do mês ÷ 2, somado nas quinzenas zeradas. É o que a placa receberia se o contrato valesse o BID; não é dívida confirmada.',
 '"Placas do mesmo modelo e unidade COM contrato" = na mesma quinzena, quantas placas iguais da mesma unidade são pagas, e o R$/km delas. Se tem irmã paga e esta não, vale olhar primeiro.',
 'Coluna Validação: preencher (ex.: contrato fabricante, carência da montadora, erro de cadastro, sem contrato).']:
    ws3.append([l])
ws3.column_dimensions['A'].width=150
wb.save('out/Placas_manutencao_zerada_FT_2026.xlsx')
print(n-1, P.faixa.value_counts().to_dict(), round(P[P.faixa=='Mais de 24 meses'].pot.sum()), P[P.faixa=='Mais de 24 meses'].irm.notna().sum())
